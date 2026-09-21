"""Word 渲染（PRD 4.1.4）：缺口章节保留并加粗标注；**被拦截引用的原文片段绝不进入文档**。

模板 v1（2-2）：新增页脚人工复核提示与页码、来源声明、模板版本号，并统一引用块
为「案号 → 原文 → 链接」三段式（见《阶段 2 第 2 阶段技术开发文档》§一 第 7 项）。
"""

from __future__ import annotations

import time
from pathlib import Path

from ..session import Session

#: 模板版本（写入文档与配置键 `export.docx_template`，便于核对交付物版本）
TEMPLATE_VERSION_FALLBACK = "research_memo_v1"
#: 来源声明（写进文档，不依赖模型自觉）
SOURCE_DECLARATION = (
    "来源声明：本报告的法条与案例引用来自北大法宝检索结果或用户上传材料，"
    "引用原文均可在「依据区」逐字核对；合成夹具数据不作为真实依据。"
)
#: 页脚固定提示（人工复核是导出的前提，必须出现在每一页）
REVIEW_NOTICE = "本报告须经人工复核后方可用于教学、研究或实务参考。"


def template_version() -> str:
    from ..config import get_config

    return str(get_config().get("export.docx_template", TEMPLATE_VERSION_FALLBACK))


def _add_page_number(paragraph) -> None:
    """在段落中插入 Word 域代码 PAGE（打开文档后自动显示当前页码）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def _build_footer(doc, version: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(REVIEW_NOTICE).bold = True
    paragraph.add_run(f"｜模板 {version}｜第 ")
    _add_page_number(paragraph)
    paragraph.add_run(" 页")


def document_text(doc) -> str:
    """汇总正文 + 表格 + 页脚的全部文本（测试与人工复核都用它做完整性断言）。"""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            parts.append(paragraph.text)
    return "\n".join(parts)


def render_research_memo(session: Session, target: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    version = template_version()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    doc.add_heading(f"类案研究报告：{session.topic or '未命名议题'}", level=0)

    meta = doc.add_paragraph()
    meta.add_run("生成工具：法小智（法律类案检索与研究）\n").bold = True
    meta.add_run(f"生成时间：{time.strftime('%Y-%m-%d %H:%M')}\n")
    meta.add_run(f"模板版本：{version}\n")
    meta.add_run(f"确认样本：{len(session.sample.confirmed)} 篇")
    if session.sample.excluded:
        meta.add_run(f"；已排除：{len(session.sample.excluded)} 篇")
    if session.gate_report is not None:
        meta.add_run(
            f"\n引用核验：通过 {session.gate_report.accepted} 条，"
            f"拦截 {len(session.gate_report.rejected)} 条"
            f"（证据覆盖率 {session.gate_report.coverage:.0%}）"
        )
    if session.gate_report is not None and session.gate_report.demo_mode:
        meta.add_run("\n⚠ 演示模式：本次引用的来源为合成夹具数据，不是真实检索结果。").bold = True
    doc.add_paragraph(SOURCE_DECLARATION)

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = notice.add_run(
        "以上分析仅供教学、学习和法律研究参考，不构成正式法律意见。"
        "涉及实际案件或重大权益，请由执业律师结合完整材料复核。"
    )
    run.bold = True

    # 材料来源清单（2-4）：本报告所用用户上传材料的角色 / 标识 / 核验 / 版本
    materials = session.material_snapshot()
    if materials:
        doc.add_heading("材料来源清单（用户上传材料）", level=1)
        columns = ["文件名", "角色", "引用标识", "核验", "版本"]
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        for index, name in enumerate(columns):
            table.rows[0].cells[index].text = name
        for record in materials:
            cells = table.add_row().cells
            identifier = str(record.get("identifier") or "")
            if record.get("identifier_missing"):
                identifier += "（未识别到案号，改用文件名）"
            cells[0].text = str(record.get("filename") or "")
            cells[1].text = str(record.get("role_label") or "")
            cells[2].text = identifier
            cells[3].text = "已核验" if record.get("verified") else (
                "已取代" if record.get("superseded") else "未核验"
            )
            cells[4].text = str(record.get("version_label") or "")
        doc.add_paragraph(
            "说明：上表材料由使用者上传，报告结论建立在这些材料之上；"
            "平台不对材料本身的真伪做鉴定，材料真实性由使用者保证（人工核验是唯一闸门）。"
        )
        # 来源构成（2-4）：把「材料为主、检索为辅」如实写进交付物。
        # 修复（代操作验收措辞不一致 a）：原只在有法宝来源时才输出，导致「纯材料路径」看不到该行；
        # 验收清单要求报告开头始终有该行，故改为无条件输出（含 0 条）。
        user_count = sum(1 for s in session.source_pool.values() if s.is_user_material)
        supplement = [s for s in session.source_pool.values() if s.supplement and not s.is_user_material]
        other = [
            s
            for s in session.source_pool.values()
            if not s.is_user_material and not s.supplement
        ]
        doc.add_paragraph(
            f"本次来源构成：用户上传材料 {user_count} 篇；"
            f"补充检索来源 {len(supplement)} 条（单独标注为「补充来源」，不改变材料的主体地位）；"
            f"常规检索来源 {len(other)} 条。用户材料在候选与矩阵中优先。"
        )

    # 人工裁决标注（2-4）：冲突被人工放行过就必须在交付物里留痕
    resolutions = list((session.conflict_resolutions or {}).values())
    if resolutions:
        paragraph = doc.add_paragraph()
        paragraph.add_run(
            "⚠ 依据用户上传材料，未经法宝印证；冲突已由人工确认"
            f"（共 {len(resolutions)} 处）。"
            + "、".join(str(r.get("identifier") or r.get("source_id")) for r in resolutions[:5])
        ).bold = True

    # 一、议题与检索条件
    doc.add_heading("一、议题与检索条件", level=1)
    doc.add_paragraph(session.topic or "（未填写）")
    if session.conditions:
        for key, value in session.conditions.items():
            if value:
                doc.add_paragraph(f"{key}：{value}", style="List Bullet")

    # 二、确认样本
    doc.add_heading("二、确认样本", level=1)
    confirmed = [
        session.source_pool[sid] for sid in session.sample.confirmed if sid in session.source_pool
    ]
    if confirmed:
        for source in confirmed:
            doc.add_paragraph(
                f"{source.identifier}｜{source.title or '（无标题）'}"
                + (f"｜{source.court}" if source.court else ""),
                style="List Bullet",
            )
    else:
        doc.add_paragraph("本次未确认任何样本。")
    if session.sample.excluded:
        excluded_note = doc.add_paragraph()
        excluded_note.add_run(
            "本次已排除（任何结论均不得引用）："
            + "、".join(session.sample.excluded)
        ).italic = True

    # 三、对比矩阵（代码计算）
    doc.add_heading("三、横向对比矩阵（由代码计算，仅含已确认样本）", level=1)
    if session.matrix:
        columns = list(session.matrix[0].keys())
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        for index, name in enumerate(columns):
            table.rows[0].cells[index].text = name
        for row in session.matrix:
            cells = table.add_row().cells
            for index, name in enumerate(columns):
                cells[index].text = str(row.get(name, ""))[:200]
    else:
        doc.add_paragraph("未生成矩阵（样本可能不足）。")

    # 四、观点分布（代码计算，永远带分母）
    doc.add_heading("四、观点分布", level=1)
    doc.add_paragraph(session.distribution.get("text") or "未生成分布。")

    # 五、综合结论
    doc.add_heading("五、综合结论（每条均带引用）", level=1)
    synthesis = session.synthesis or {}
    passed_texts = (session.structured_output or {}).get("_passed_conclusions", [])
    conclusions = synthesis.get("conclusions") or []
    if conclusions:
        for index, item in enumerate(conclusions, start=1):
            text = str(item.get("text", ""))
            if text not in passed_texts:
                continue
            doc.add_paragraph(f"{index}. {text}")
            for citation in item.get("citations") or []:
                source = session.source_pool.get(str(citation.get("source_id")))
                # v1 统一引用块：案号（加粗）→ 原文 → 链接；缺失的段落不写空壳
                line = doc.add_paragraph(style="List Bullet")
                line.add_run(f"引用：{citation.get('identifier', '')}").bold = True
                quote = str(citation.get("quote", ""))
                if quote:
                    line.add_run(f"｜原文：「{quote}」")
                if source is not None and source.uri:
                    line.add_run(f"｜链接：{source.uri}")
                if source is not None:
                    origin_text = source.origin_text
                    line.add_run(f"｜来源：{origin_text}｜印证：{source.corroboration_text}")
                    if source.manual_override:
                        line.add_run("｜已由人工裁决采用用户材料（未经法宝印证）")
                    if source.local_hit:
                        line.add_run("｜本地命中")
    else:
        doc.add_paragraph("本次未产出综合结论。")

    # 六、依据缺口（保留并加粗标注；不写入被拦截的原文片段）
    doc.add_heading("六、依据缺口（未作为结论输出的部分）", level=1)
    if not session.gaps:
        doc.add_paragraph("无。")
    else:
        for gap in session.gaps:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(f"[{gap.kind}] {gap.detail}").bold = True
    if session.gate_report and session.gate_report.rejected:
        doc.add_paragraph("未通过核验的引用（仅列原因，原文片段不予收录）：")
        for rejection in session.gate_report.rejected:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{rejection.rule}：{rejection.reason}").bold = True

    # 七、方法与局限
    doc.add_heading("七、方法与局限", level=1)
    doc.add_paragraph(
        "本报告基于用户确认的样本生成，横向对比矩阵与观点分布均由程序计算，"
        "不包含未确认样本，也不代表全国裁判总体情况。"
    )
    for limitation in session.limitations:
        doc.add_paragraph(limitation, style="List Bullet")

    footer = doc.add_paragraph()
    footer.add_run(REVIEW_NOTICE).bold = True

    # v1：页脚（每页可见）——人工复核提示 + 模板版本 + 页码
    _build_footer(doc, version)

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target


#: 咨询备忘声明（2-5）：与研究报告同一套口径，额外强调「不构成正式法律意见」
CONSULT_DECLARATION = (
    "本备忘基于检索到的法条与案例生成，引用原文均可在「依据区」逐字核对；"
    "仅作教学、学习与研究参考，**不构成正式法律意见，也不可替代执业律师**。"
    "涉及实际案件或重大权益，请由执业律师结合完整材料复核。"
)


def render_consult_memo(session: Session, target: Path) -> Path:
    """渲染「咨询备忘」（阶段 2-5）。

    与研究报告共用同一套底线：**被拦截引用的原文片段绝不进入文档**；
    只有通过引用门禁的结论才会出现在「结论」章；检索不到依据时如实写出来。
    """
    from docx import Document
    from docx.shared import Pt

    version = template_version()
    state = session.consult or {}
    answer = state.get("answer") or {}
    issue = state.get("issue") or {}
    passed_texts = (session.structured_output or {}).get("_passed_conclusions") or []
    report = session.gate_report

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    doc.add_heading(f"咨询备忘：{session.topic or '未命名问题'}", level=0)

    meta = doc.add_paragraph()
    meta.add_run("生成工具：法小智（法律咨询）\n").bold = True
    meta.add_run(f"生成时间：{time.strftime('%Y-%m-%d %H:%M')}\n")
    meta.add_run(f"模板版本：{version}\n")
    if report is not None:
        meta.add_run(
            f"引用核验：通过 {report.accepted} 条 / 拦截 {len(report.rejected)} 条"
            f"（证据覆盖率 {report.coverage * 100:.0f}%）\n"
        )
    meta.add_run(
        f"追问轮数：{state.get('rounds', 0)} / {state.get('max_rounds', 3)}\n"
    )
    doc.add_paragraph(SOURCE_DECLARATION)
    doc.add_paragraph(CONSULT_DECLARATION)

    # 一、问题识别
    doc.add_heading("一、问题识别（由代码记录，非结论）", level=1)
    doc.add_paragraph(f"法律关系：{issue.get('legal_relation') or '（未识别）'}")
    for label, key in (
        ("争议焦点", "disputes"),
        ("已给事实", "given_facts"),
        ("缺失事实", "missing_facts"),
        ("用户可能的主张", "possible_claims"),
    ):
        values = [str(v) for v in (issue.get(key) or []) if str(v).strip()]
        if not values:
            continue
        doc.add_paragraph(f"{label}：", style="List Bullet")
        for value in values:
            doc.add_paragraph(f"    - {value}", style="List Bullet 2")

    # 二、结论（只有通过门禁的才写出）
    doc.add_heading("二、结论（每条均带可回查引用）", level=1)
    conclusions = answer.get("conclusions") or []
    rendered = 0
    for item in conclusions:
        text = str(item.get("text", ""))
        if text not in passed_texts:
            continue
        rendered += 1
        doc.add_paragraph(f"{rendered}. {text}")
        _render_citations(doc, session, item.get("citations") or [])
    if rendered == 0:
        insufficient = str(answer.get("insufficient") or "").strip()
        doc.add_paragraph(
            insufficient
            or "本次未产出可引用依据支撑的结论（检索无匹配或引用未通过核验）。"
        )

    # 三、依据清单
    doc.add_heading("三、本次用到的依据", level=1)
    cited_ids: list[str] = []
    for item in conclusions:
        for citation in item.get("citations") or []:
            sid = str(citation.get("source_id") or "")
            if sid and sid not in cited_ids:
                cited_ids.append(sid)
    if cited_ids:
        for sid in cited_ids:
            source = session.source_pool.get(sid)
            if source is None:
                continue
            line = doc.add_paragraph(style="List Bullet")
            line.add_run(f"{source.identifier}").bold = True
            line.add_run(f"｜{source.title or '（无标题）'}")
            line.add_run(f"｜来源：{source.origin_text}｜印证：{source.corroboration_text}")
    else:
        doc.add_paragraph("本次没有可引用的依据。")

    # 四、不确定与风险
    doc.add_heading("四、不确定与风险", level=1)
    uncertainties = [str(x) for x in (answer.get("uncertainties") or []) if str(x).strip()]
    if uncertainties:
        for item in uncertainties:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("（本次未列出不确定项——这是异常情况，请人工复核。）")

    # 五、建议的下一步
    doc.add_heading("五、建议的下一步", level=1)
    steps = [str(x) for x in (answer.get("next_steps") or []) if str(x).strip()]
    if steps:
        for item in steps:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("（未给出建议。）")

    # 六、依据缺口（未被采用的引用不写原文）
    doc.add_heading("六、依据缺口（未作为结论输出的部分）", level=1)
    if not session.gaps:
        doc.add_paragraph("本次没有记录到依据缺口。")
    else:
        for gap in session.gaps:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(f"[{gap.kind}] ").bold = True
            paragraph.add_run(gap.detail)
    if report is not None and report.rejected:
        paragraph = doc.add_paragraph()
        paragraph.add_run(
            f"另有 {len(report.rejected)} 条引用未通过核验，相关结论已降级，未写入本备忘。"
        ).bold = True

    # 七、方法与局限
    doc.add_heading("七、方法与局限", level=1)
    doc.add_paragraph(
        "本备忘基于本次检索到的法条与案例生成；引用核验、表达边界与追问轮数均由程序判定，"
        "不代表全国司法裁判总体情况，也不构成对结果的承诺。"
    )
    for limitation in session.limitations:
        doc.add_paragraph(limitation, style="List Bullet")

    footer = doc.add_paragraph()
    footer.add_run("本备忘须经人工复核后方可用于教学、研究或实务参考。").bold = True
    _build_footer(doc, version)

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target


def _render_citations(doc, session: Session, citations: list) -> None:
    """统一引用块（与研究报告 v1 同格式）：案号 → 原文 → 链接 → 来源与印证。"""
    for citation in citations:
        source = session.source_pool.get(str(citation.get("source_id")))
        line = doc.add_paragraph(style="List Bullet")
        line.add_run(f"引用：{citation.get('identifier', '')}").bold = True
        quote = str(citation.get("quote", ""))
        if quote:
            line.add_run(f"｜原文：「{quote}」")
        if source is not None and source.uri:
            line.add_run(f"｜链接：{source.uri}")
        if source is not None:
            line.add_run(f"｜来源：{source.origin_text}｜印证：{source.corroboration_text}")
            if source.manual_override:
                line.add_run("｜已由人工裁决采用用户材料（未经法宝印证）")
            if source.local_hit:
                line.add_run("｜本地命中")


#: 合同审查声明（2-6）：在研究报告口径上再加一条「未提示不等于无风险」
CONTRACT_DECLARATION = (
    "本报告由法小智（合同审查）生成，法律依据来自北大法宝检索结果或用户上传并经核验的材料，"
    "引用原文均可在「依据区」逐字核对。"
    "**本报告不构成正式法律意见，也不可替代执业律师**；建议改法仅为示例表述，须经律师审定。"
    "**未被提示风险的条款不等于没有风险**：本报告基于当前检索到的依据与模型判断，可能存在漏报。"
)


def render_contract_report(session: Session, target: Path, review: dict | None = None) -> Path:
    """渲染《合同审查报告》（阶段 2-6）。

    与其它报告共用同一套底线：**被拦截引用的原文片段绝不进入文档**；
    只有通过引用门禁的依据才作为「法律依据」展示，其它降级为「提示性风险（未找到直接依据）」。
    """
    from docx import Document
    from docx.shared import Pt

    from ..contract import RISK_KIND_LABELS, RISK_LEVEL_LABELS

    version = template_version()
    state = session.contract or {}
    review = review or {}
    risks = list(state.get("risks") or [])
    counts = state.get("counts") or {}
    report = session.gate_report

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    doc.add_heading(f"合同审查报告：{state.get('filename') or '未命名合同'}", level=0)
    meta = doc.add_paragraph()
    meta.add_run("生成工具：法小智（合同审查）\n").bold = True
    meta.add_run(f"生成时间：{time.strftime('%Y-%m-%d %H:%M')}\n")
    meta.add_run(f"模板版本：{version}\n")
    meta.add_run(f"审查立场：{state.get('stanceLabel') or state.get('stance') or '（未确认）'}\n")
    if report is not None:
        meta.add_run(
            f"引用核验：通过 {report.accepted} 条 / 拦截 {len(report.rejected)} 条"
            f"（证据覆盖率 {report.coverage * 100:.0f}%）\n"
        )
    meta.add_run(
        f"风险条目：{counts.get('total', len(risks))} 条"
        f"（有依据 {counts.get('verified', 0)}｜提示性 {counts.get('no_basis', 0)}"
        f"｜定位失败 {counts.get('locate_failed', 0)}）\n"
    )
    doc.add_paragraph(SOURCE_DECLARATION)
    doc.add_paragraph(CONTRACT_DECLARATION)

    # 一、合同基本信息
    doc.add_heading("一、合同基本信息", level=1)
    parties = state.get("parties") or {}
    doc.add_paragraph(f"文件：{state.get('filename') or '（未命名）'}")
    doc.add_paragraph(f"条款数：{len(state.get('clauses') or [])}")
    if parties:
        doc.add_paragraph(f"甲方：{parties.get('party_a') or '（未识别）'}")
        doc.add_paragraph(f"乙方：{parties.get('party_b') or '（未识别）'}")
    doc.add_paragraph(f"审查立场：{state.get('stanceLabel') or state.get('stance') or '（未确认）'}")

    # 二、风险清单（按等级排序：高 → 中 → 低）
    doc.add_heading("二、风险清单（按等级排序）", level=1)
    order = {"high": 0, "medium": 1, "low": 2}
    ranked = sorted(risks, key=lambda r: order.get(str(r.get("level")), 9))
    if not ranked:
        doc.add_paragraph("本次未识别到风险条目。**这不等于合同没有风险**，请人工复核。")
    for index, risk in enumerate(ranked, start=1):
        heading = doc.add_paragraph()
        heading.add_run(
            f"{index}. [{RISK_LEVEL_LABELS.get(str(risk.get('level')), risk.get('level'))}] "
            f"{RISK_KIND_LABELS.get(str(risk.get('kind')), risk.get('kind'))}"
            f"　{risk.get('clauseHeading') or risk.get('clauseId') or ''}"
        ).bold = True
        doc.add_paragraph(f"风险说明：{risk.get('issue') or ''}")
        anchor = str(risk.get("anchorText") or "")
        if risk.get("locateOk") and anchor:
            paragraph = doc.add_paragraph()
            paragraph.add_run("原文位置：").bold = True
            paragraph.add_run(f"第 {risk.get('clauseId')} 块｜「{anchor}」")
        else:
            paragraph = doc.add_paragraph()
            paragraph.add_run("原文位置：").bold = True
            paragraph.add_run("未能定位到合同原文（锚点对不上），请人工核对。")
        status = risk.get("basisStatus")
        if risk.get("basis") and status == "verified":
            _render_citations(doc, session, risk.get("basis") or [])
        elif risk.get("basis"):
            paragraph = doc.add_paragraph()
            paragraph.add_run(
                "⚠ 该条给出的法律依据未通过引用核验，已降级为「提示性风险」，不作为有依据的风险。"
            ).bold = True
        else:
            paragraph = doc.add_paragraph()
            paragraph.add_run("⚠ 未找到直接依据（提示性风险，非无风险）。").bold = True
        if risk.get("suggestion"):
            doc.add_paragraph(f"建议改法（示例表述，须经律师审定）：{risk.get('suggestion')}")

    if review:
        doc.add_heading("人工处理记录（未进行自动法律核验）", level=1)
        doc.add_paragraph("以下为用户对本次审查的处理及编辑意见，不改变系统的依据核验状态，也未修改合同原文件。")
        for risk in ranked:
            note = review.get(str(risk.get("riskId")), {})
            label = {"accept": "已采纳", "skip": "暂不采纳", "pending": "待处理"}.get(note.get("status"), "待处理")
            doc.add_paragraph(f"{risk.get('clauseHeading') or risk.get('riskId')}：{label}")
            if note.get("suggestion"):
                doc.add_paragraph(f"人工处理意见（须人工复核）：{note['suggestion']}")

    # 三、依据缺口
    doc.add_heading("三、依据缺口（未作为有依据风险输出的部分）", level=1)
    if not session.gaps:
        doc.add_paragraph("本次没有记录到依据缺口。")
    else:
        for gap in session.gaps:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(f"[{gap.kind}] ").bold = True
            paragraph.add_run(gap.detail)
    if report is not None and report.rejected:
        paragraph = doc.add_paragraph()
        paragraph.add_run(
            f"另有 {len(report.rejected)} 条引用未通过核验，未写入本报告。"
        ).bold = True

    # 四、方法与局限
    doc.add_heading("四、方法与局限", level=1)
    doc.add_paragraph(
        "条款切分与原文位置由程序计算；风险识别与等级由模型给出、取值范围由程序校验；"
        "引用核验、表达边界与上限截断均由程序判定。"
        "**未提示风险的条款不等于没有风险**，本报告可能存在漏报。"
    )
    for limitation in session.limitations:
        doc.add_paragraph(limitation, style="List Bullet")

    footer = doc.add_paragraph()
    footer.add_run("本报告须经人工复核后方可用于教学、研究或实务参考。").bold = True
    _build_footer(doc, version)

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target
